"""Formulario 1 del IDU en miniatura, con todos los casos del archivo real.

Reproduce: membrete, encabezado en una fila que no es la 1, dos capítulos, turno diurno
y nocturno, subtítulos, filas Subtotal, un encabezado repetido a mitad, filas de resumen
global, un ítem de pago con cero significativo (guardado como float con formato 0.000,
igual que Excel) y otro con coma y sufijo -N.

Los valores están elegidos para que round(cantidad * con_aiu) dé exacto, igual que en el
archivo real (1939 de 1939 filas).
"""
import openpyxl

# Estructura del archivo real (columnas 0-based):
# 2=Nº  3=ITEM DE PAGO  6=DESCRIPCION  7=UND.  8=CANTIDAD
# 9=VALOR UNITARIO BASICO (SIN A.I.U)  10=VALOR UNITARIO (INCLUYE A.I.U)  11=VALOR TOTAL
COL_CODIGO, COL_ITEM, COL_DESC, COL_UND, COL_CANT = 2, 3, 6, 7, 8
COL_SIN_AIU, COL_CON_AIU, COL_TOTAL = 9, 10, 11

ENCABEZADO = {
    COL_CODIGO: "Nº", COL_ITEM: "ITEM DE PAGO", 4: "ESPECIFICACIONES ",
    COL_DESC: "DESCRIPCION", COL_UND: "UND.", COL_CANT: "CANTIDAD",
    COL_SIN_AIU: "VALOR UNITARIO BASICO (SIN A.I.U)",
    COL_CON_AIU: "VALOR UNITARIO (INCLUYE A.I.U)",
    COL_TOTAL: "VALOR TOTAL                    ",
    14: "VALOR UNITARIO SIN AIU OFERTADO",
}

# (código, ítem de pago, formato del ítem, descripción, und, cantidad, sin AIU, con AIU)
ACTIVIDADES_CAP_1 = [
    (3007, 1.001, "0.000", "REPLANTEO GENERAL", "M2", 100, 1056, 1351),
    ("3007 N", "1,001-N", "General", "REPLANTEO GENERAL", "M2", 40, 1179, 1508),
]
ACTIVIDADES_CAP_2 = [
    (3710, 2.001, "0.000", "EXCAVACIÓN MECÁNICA", "M3", 30, 5963, 7628),
    # El cero significativo: Excel guarda 2.010 como el float 2.01 con formato 0.000.
    (3866, 2.01, "0.000", "RIEGO DE LIGA", "M2", 50, 3815, 4880),
    ("4200 N", "2,011-N", "General", "MEZCLA ASFÁLTICA MD20", "M3", 7, 1188814, 1520731),
]


def total_esperado(cant, con_aiu) -> int:
    return round(cant * con_aiu)


def escribir_formulario(path, *, hoja="PROPUESTA ECONÓMICA", con_hoja_gemela=False,
                        sin_columna_cantidad=False):
    """Escribe el Formulario 1 en miniatura y devuelve la ruta.

    `con_hoja_gemela` agrega una segunda hoja candidata (el caso real de las dos
    PROPUESTA ECONÓMICA). `sin_columna_cantidad` borra un encabezado obligatorio para
    probar el error bloqueante.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = hoja
    _volcar(ws, sin_columna_cantidad=sin_columna_cantidad)
    if con_hoja_gemela:
        _volcar(wb.create_sheet(f"{hoja} (2)"), sin_columna_cantidad=sin_columna_cantidad)
    wb.create_sheet("INDICAR CÓDIGO DEL ITEM DE PAGO")
    wb.save(path)
    return path


def _volcar(ws, *, sin_columna_cantidad=False):
    ancho = 16

    def fila(cols=None, formatos=None):
        ws.append([(cols or {}).get(i) for i in range(ancho)])
        for idx, fmt in (formatos or {}).items():
            ws.cell(row=ws.max_row, column=idx + 1).number_format = fmt

    fila()                                            # 1: vacía
    fila({5: "FORMULARIO 1"})                         # 2: membrete
    fila({5: "PRESUPUESTO OFICIAL ESTIMADO"})         # 3: membrete
    fila()                                            # 4
    enc = dict(ENCABEZADO)
    if sin_columna_cantidad:
        enc.pop(COL_CANT)
    fila(enc)                                         # 5: ENCABEZADO
    fila({4: "GENERAL", 5: "PARTICULAR"})             # 6: subencabezado, se ignora
    _capitulo(fila, 1, "PRELIMINARES", ACTIVIDADES_CAP_1, "LOCALIZACIÓN Y REPLANTEO")
    _capitulo(fila, 2, "PAVIMENTOS", ACTIVIDADES_CAP_2, "PAVIMENTO FLEXIBLE")
    fila()
    fila(dict(ENCABEZADO))                            # encabezado REPETIDO a mitad
    fila({4: "GENERAL", 5: "PARTICULAR"})
    fila()
    fila({COL_CODIGO: "VALOR PARA OBRAS SIN REDES (INCLUYE A.I.U)", COL_TOTAL: 999999})
    fila({COL_CODIGO: "TOTAL OBRAS LICITACIÓN ( A+B+C+D)", COL_TOTAL: 999999})
    fila()
    fila({0: "Total general"})                        # cola del archivo real


def _capitulo(fila, numero, nombre, actividades, subtitulo):
    fila({COL_ITEM: numero, COL_DESC: nombre}, {COL_ITEM: "0"})
    subtotal = 0
    for i, (cod, item, fmt, desc, und, cant, sin_aiu, con_aiu) in enumerate(actividades):
        if i == 0:
            fila({COL_DESC: "TURNO DIURNO"})
            fila({COL_DESC: subtitulo})
        elif i == 1:
            fila({COL_DESC: "TURNO NOCTURNO"})
            fila({COL_DESC: subtitulo})
        total = total_esperado(cant, con_aiu)
        subtotal += total
        fila({COL_CODIGO: cod, COL_ITEM: item, COL_DESC: desc, COL_UND: und,
              COL_CANT: cant, COL_SIN_AIU: sin_aiu, COL_CON_AIU: con_aiu,
              COL_TOTAL: total},
             {COL_ITEM: fmt})
    fila({COL_CODIGO: "Subtotal ", COL_TOTAL: subtotal})
