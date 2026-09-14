"""Parser del Formulario 1 de Presupuesto Oficial del IDU."""
from apu_tool.dominio.presupuesto import (
    capitulo_de, es_item_entero, item_pago_texto, norm_encabezado, normalizar_item_pago,
)


def test_norm_encabezado_quita_tildes_y_mayusculas():
    assert norm_encabezado("DESCRIPCION") == "descripcion"
    assert norm_encabezado("DESCRIPCIÓN") == "descripcion"
    assert norm_encabezado("ÍTEM DE PAGO") == "item de pago"
    assert norm_encabezado("ITEM DE PAGO") == "item de pago"


def test_norm_encabezado_colapsa_espacios_y_saltos():
    assert norm_encabezado("VALOR TOTAL                 ") == "valor total"
    assert norm_encabezado("VALOR UNITARIO SIN AIU\nCORREGIDO") == \
        "valor unitario sin aiu corregido"
    assert norm_encabezado("  CANTIDAD  ") == "cantidad"


def test_norm_encabezado_unifica_las_variantes_de_numero():
    # Nº (ordinal masculino), N° (grado), No. y No son el mismo encabezado.
    assert norm_encabezado("Nº") == "no"
    assert norm_encabezado("N°") == "no"
    assert norm_encabezado("No.") == "no"
    assert norm_encabezado("No") == "no"


def test_norm_encabezado_quita_puntos_y_parentesis():
    assert norm_encabezado("UND.") == "und"
    assert norm_encabezado("VALOR UNITARIO BASICO (SIN A.I.U)") == \
        "valor unitario basico sin aiu"
    assert norm_encabezado("VALOR UNITARIO (INCLUYE A.I.U)") == \
        "valor unitario incluye aiu"


def test_norm_encabezado_tolera_none_y_numeros():
    assert norm_encabezado(None) == ""
    assert norm_encabezado(3007) == "3007"
    # openpyxl entrega los enteros de una celda como float: 3007.0, no 3007. Sin el
    # guard, el punto decimal se pierde como puntuación y esto daría "30070".
    assert norm_encabezado(3007.0) == "3007"


def test_item_pago_texto_recupera_el_cero_perdido_por_el_float():
    # Excel guarda 2.010 como el float 2.01; el formato de celda dice cuántos
    # decimales tenía. Sin esto, el ítem de pago del pliego no se puede cruzar.
    assert item_pago_texto(2.01, "0.000") == "2.010"
    assert item_pago_texto(3.1, "0.000") == "3.100"
    assert item_pago_texto(2.001, "0.000") == "2.001"


def test_item_pago_texto_capitulo_es_entero():
    assert item_pago_texto(2, "0") == "2"
    assert item_pago_texto(14.0, "0") == "14"


def test_item_pago_texto_pasa_el_texto_tal_cual():
    assert item_pago_texto("1,001-N", "General") == "1,001-N"
    assert item_pago_texto("  2.014 N  ", "General") == "2.014 N"
    assert item_pago_texto(None, "0.000") == ""


def test_item_pago_texto_sin_formato_util_no_inventa_decimales():
    assert item_pago_texto(2.001, "General") == "2.001"
    assert item_pago_texto(7.0, "General") == "7"
    # Sin formato útil tampoco puede PERDER dígitos: un `:g` cortaría a 6 cifras
    # significativas y esto daría "123457".
    assert item_pago_texto(123456.789, "General") == "123456.789"


def test_item_pago_texto_descarta_booleanos_y_nan():
    # `bool` es `int` en Python: sin el guard, True se leería como el ítem "1" y la
    # fila entraría al capítulo 1. NaN es lo que deja una fórmula rota.
    assert item_pago_texto(True, "0.000") == ""
    assert item_pago_texto(False, "0.000") == ""
    assert item_pago_texto(float("nan"), "0.000") == ""


def test_normalizar_item_pago_unifica_separadores_y_sufijo_de_turno():
    assert normalizar_item_pago("2.001") == "2.001"
    assert normalizar_item_pago("2,001") == "2.001"
    assert normalizar_item_pago("2.001 N") == "2.001 N"
    assert normalizar_item_pago("2.001-N") == "2.001 N"
    assert normalizar_item_pago("2,001-N") == "2.001 N"
    assert normalizar_item_pago('"2.001"') == "2.001"
    assert normalizar_item_pago("'2.001'") == "2.001"
    # Comillas anidadas al revés: la simple por fuera, la doble por dentro.
    assert normalizar_item_pago("'\"2.001\"'") == "2.001"


def test_normalizar_item_pago_no_pierde_ceros():
    # La normalización es TEXTO: nunca pasa por float, así que el cero sobrevive.
    assert normalizar_item_pago("2.010") == "2.010"
    assert normalizar_item_pago("3.100") == "3.100"


def test_capitulo_de_saca_el_prefijo_entero():
    assert capitulo_de("2.014") == "2"
    assert capitulo_de("2.014 N") == "2"
    assert capitulo_de("2,014-N") == "2"
    assert capitulo_de("14.057-N") == "14"
    assert capitulo_de("2") == "2"
    assert capitulo_de("") == ""
    assert capitulo_de("SUBTOTAL") == ""
    # Celda vacía de Excel: las tareas siguientes llaman esto sobre valores crudos.
    assert capitulo_de(None) == ""


def test_es_item_entero_distingue_capitulo_de_actividad():
    assert es_item_entero("2") is True
    assert es_item_entero("14") is True
    assert es_item_entero("2.001") is False
    assert es_item_entero("2.001 N") is False
    assert es_item_entero("") is False
    assert es_item_entero(None) is False


# --------------------------------------------------------------- clasificación
from apu_tool.dominio.presupuesto import (  # noqa: E402
    ACTIVIDAD, CAPITULO, ENCABEZADO, IGNORADA, SUBTITULO, SUBTOTAL, TURNO,
    clasificar_fila,
)


def _clasificar(codigo="", item_pago="", descripcion="", cantidad=None,
                es_encabezado=False):
    return clasificar_fila(codigo=codigo, item_pago=item_pago,
                           descripcion=descripcion, cantidad=cantidad,
                           es_encabezado=es_encabezado)


def test_clasificar_capitulo():
    # Fila 13 del archivo real: Nº vacío, ítem de pago entero, descripción.
    assert _clasificar(item_pago="1", descripcion="PRELIMINARES") == CAPITULO
    assert _clasificar(item_pago="10", descripcion="RED DE ALCANTARILLADO") == CAPITULO


def test_clasificar_actividad():
    # Fila 16: código IDU + cantidad > 0.
    assert _clasificar(codigo="3007", item_pago="1.001",
                       descripcion="REPLANTEO GENERAL", cantidad=74234) == ACTIVIDAD
    # Nocturna: el código trae el sufijo N.
    assert _clasificar(codigo="3007 N", item_pago="1.001 N",
                       descripcion="REPLANTEO GENERAL", cantidad=24745) == ACTIVIDAD


def test_clasificar_turno_subtitulo_y_subtotal():
    assert _clasificar(descripcion="TURNO DIURNO") == TURNO
    assert _clasificar(descripcion="TURNO NOCTURNO") == TURNO
    assert _clasificar(descripcion="LOCALIZACIÓN Y REPLANTEO") == SUBTITULO
    assert _clasificar(descripcion="PAVIMENTO RÍGIDO") == SUBTITULO
    assert _clasificar(codigo="Subtotal ") == SUBTOTAL
    assert _clasificar(codigo="Subtotal") == SUBTOTAL


def test_clasificar_encabezado_repetido():
    # Fila 2201: el segundo encabezado NO puede contar como actividad.
    assert _clasificar(codigo="Nº", item_pago="ÍTEM DE PAGO",
                       descripcion="DESCRIPCIÓN", es_encabezado=True) == ENCABEZADO


def test_clasificar_resumen_notas_y_vacias():
    # Filas 2204+: texto largo en la columna Nº, sin descripción y sin cantidad.
    assert _clasificar(codigo="VALOR PARA OBRAS SIN REDES (INCLUYE A.I.U)") == IGNORADA
    assert _clasificar(codigo="TOTAL OBRAS LICITACIÓN ( A+B+C+D)") == IGNORADA
    assert _clasificar() == IGNORADA


def test_clasificar_no_cuenta_actividad_sin_cantidad_positiva():
    assert _clasificar(codigo="3007", item_pago="1.001", descripcion="X",
                       cantidad=0) == IGNORADA
    assert _clasificar(codigo="3007", item_pago="1.001", descripcion="X",
                       cantidad=None) == IGNORADA
    assert _clasificar(codigo="3007", item_pago="1.001", descripcion="X",
                       cantidad="n/a") == IGNORADA


def test_clasificar_capitulo_gana_sobre_subtitulo():
    # Con ítem entero manda CAPITULO aunque también tenga descripción.
    assert _clasificar(item_pago="2", descripcion="PAVIMENTOS") == CAPITULO
    # Sin ítem entero, la misma forma es un subtítulo.
    assert _clasificar(item_pago="", descripcion="PAVIMENTOS") == SUBTITULO


# ------------------------------------------- detección de hoja y de encabezado
from apu_tool.dominio.presupuesto import (  # noqa: E402
    COLUMNAS_OBLIGATORIAS, _es_fila_encabezado, elegir_hoja,
    encontrar_encabezado,
)


class _LibroFalso:
    """Stand-in de un Workbook: `elegir_hoja` solo necesita `sheetnames`."""
    def __init__(self, nombres):
        self.sheetnames = list(nombres)


def test_elegir_hoja_encuentra_propuesta_economica():
    wb = _LibroFalso(["PROPUESTA ECONÓMICA", "INDICAR CÓDIGO DEL ITEM DE PAGO"])
    hoja, ambiguas = elegir_hoja(wb, None)
    assert hoja == "PROPUESTA ECONÓMICA"
    assert ambiguas == []


def test_elegir_hoja_avisa_cuando_hay_varias_candidatas():
    # El archivo real trae dos hojas idénticas: se usa la primera y se avisa.
    wb = _LibroFalso(["PROPUESTA ECONÓMICA", "PROPUESTA ECONÓMICA (2)", "OTRA"])
    hoja, ambiguas = elegir_hoja(wb, None)
    assert hoja == "PROPUESTA ECONÓMICA"
    assert ambiguas == ["PROPUESTA ECONÓMICA (2)"]


def test_elegir_hoja_acepta_las_variantes_conocidas():
    assert elegir_hoja(_LibroFalso(["FOR 1-PPTO OFICIAL"]), None)[0] == "FOR 1-PPTO OFICIAL"
    assert elegir_hoja(_LibroFalso(["Presupuesto Oficial"]), None)[0] == "Presupuesto Oficial"


def test_elegir_hoja_respeta_la_hoja_explicita():
    wb = _LibroFalso(["PROPUESTA ECONÓMICA", "HOJA RARA"])
    assert elegir_hoja(wb, "HOJA RARA")[0] == "HOJA RARA"


def test_elegir_hoja_sin_candidata_devuelve_none():
    hoja, ambiguas = elegir_hoja(_LibroFalso(["Hoja1", "Datos"]), None)
    assert hoja is None
    assert ambiguas == []


def test_elegir_hoja_explicita_inexistente_devuelve_none():
    assert elegir_hoja(_LibroFalso(["PROPUESTA ECONÓMICA"]), "NO EXISTE")[0] is None


_ENCABEZADO_REAL = [
    None, None, "Nº", "ITEM DE PAGO", "ESPECIFICACIONES ", None, "DESCRIPCION",
    "UND.", "CANTIDAD", "VALOR UNITARIO BASICO (SIN A.I.U)",
    "VALOR UNITARIO (INCLUYE A.I.U)", "VALOR TOTAL                    ",
    None, None, "VALOR UNITARIO SIN AIU OFERTADO",
]


def test_encontrar_encabezado_mapea_las_columnas_por_nombre():
    filas = [[None] * 15, [None] * 15, _ENCABEZADO_REAL, [None] * 15]
    idx, mapeo, faltan = encontrar_encabezado(filas)
    assert idx == 2                       # 0-based; es la fila 3 del Excel
    assert faltan == []
    assert mapeo["codigo"] == 2
    assert mapeo["item_pago"] == 3
    assert mapeo["descripcion"] == 6
    assert mapeo["unidad"] == 7
    assert mapeo["cantidad"] == 8
    assert mapeo["unitario_sin_aiu"] == 9
    assert mapeo["unitario_con_aiu"] == 10
    assert mapeo["total_excel"] == 11


def test_encontrar_encabezado_ignora_las_columnas_de_oferta():
    # `VALOR UNITARIO SIN AIU OFERTADO` no debe robarle el mapeo a la col J.
    _idx, mapeo, _faltan = encontrar_encabezado([_ENCABEZADO_REAL])
    assert mapeo["unitario_sin_aiu"] == 9
    assert 14 not in mapeo.values()


def test_encontrar_encabezado_reporta_las_columnas_que_faltan():
    fila = [None, None, "Nº", "ITEM DE PAGO", None, None, "DESCRIPCION", None, None]
    idx, _mapeo, faltan = encontrar_encabezado([fila])
    assert idx == 0
    assert "cantidad" in faltan
    assert "unitario_con_aiu" in faltan


def test_encontrar_encabezado_sin_encabezado_devuelve_menos_uno():
    idx, mapeo, faltan = encontrar_encabezado([[None] * 8, ["a", "b", "c"]])
    assert idx == -1
    assert mapeo == {}
    assert faltan == sorted(COLUMNAS_OBLIGATORIAS)


# ------------------------------------------------- fixture del Formulario 1
import openpyxl  # noqa: E402

from tests.fixtures_idu import (  # noqa: E402
    ACTIVIDADES_CAP_1, ACTIVIDADES_CAP_2, escribir_formulario,
)


def test_fixture_conserva_el_formato_del_item_de_pago(tmp_path):
    # Si el fixture no reprodujera el float CON su formato, el test del cero
    # significativo estaría probando algo que el archivo real no hace.
    p = escribir_formulario(tmp_path / "f1.xlsx")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb["PROPUESTA ECONÓMICA"]
    formatos = {c.value: c.number_format
                for fila in ws.iter_rows(min_col=4, max_col=4) for c in fila
                if isinstance(c.value, float)}
    wb.close()
    assert formatos[2.01] == "0.000"


def test_fixture_reproduce_la_estructura_del_archivo_real(tmp_path):
    # El encabezado NO está en la fila 1, y hay un segundo encabezado a mitad.
    p = escribir_formulario(tmp_path / "f1.xlsx")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    filas = [[c.value for c in f] for f in wb["PROPUESTA ECONÓMICA"].iter_rows()]
    wb.close()
    idx, mapeo, faltan = encontrar_encabezado(filas)
    assert idx == 4                      # 0-based; fila 5 del Excel
    assert faltan == []
    assert _es_fila_encabezado(filas[24], mapeo)   # el encabezado repetido
    n_actividades = len(ACTIVIDADES_CAP_1) + len(ACTIVIDADES_CAP_2)
    assert n_actividades == 5


# ----------------------------------------------------- el lector completo
from apu_tool.dominio.presupuesto import leer_formulario_idu  # noqa: E402


def test_lee_hoja_encabezado_y_estructura(tmp_path):
    lec = leer_formulario_idu(escribir_formulario(tmp_path / "f1.xlsx"))
    assert lec.errores == []
    assert lec.hoja == "PROPUESTA ECONÓMICA"
    assert lec.fila_encabezado == 5               # 1-based, como lo ve el usuario
    assert lec.parser_version == "idu-f1/1"
    assert [c.codigo for c in lec.capitulos] == ["1", "2"]
    assert [c.nombre for c in lec.capitulos] == ["PRELIMINARES", "PAVIMENTOS"]
    assert [c.orden for c in lec.capitulos] == [1, 2]
    assert len(lec.items) == len(ACTIVIDADES_CAP_1) + len(ACTIVIDADES_CAP_2)


def test_no_cuenta_turnos_subtitulos_subtotales_ni_encabezados(tmp_path):
    lec = leer_formulario_idu(escribir_formulario(tmp_path / "f1.xlsx"))
    descripciones = {i.descripcion for i in lec.items}
    for basura in ("TURNO DIURNO", "TURNO NOCTURNO", "PRELIMINARES", "PAVIMENTOS",
                   "LOCALIZACIÓN Y REPLANTEO", "PAVIMENTO FLEXIBLE", "DESCRIPCION",
                   "VALOR PARA OBRAS SIN REDES (INCLUYE A.I.U)"):
        assert basura not in descripciones
    assert lec.filas_ignoradas > 0
    assert any(a.tipo == "encabezado_repetido" for a in lec.advertencias)


def test_cada_actividad_hereda_capitulo_y_turno(tmp_path):
    lec = leer_formulario_idu(escribir_formulario(tmp_path / "f1.xlsx"))
    por_item = {i.item_pago_original: i for i in lec.items}
    diurna = por_item["1.001"]
    assert diurna.capitulo_codigo == "1"
    assert diurna.capitulo_nombre == "PRELIMINARES"
    assert diurna.categoria == "1 · PRELIMINARES"   # derivado, para report_categorizado
    assert diurna.shift == "DIURNO"
    assert diurna.codigo_sugerido == "3007"
    nocturna = por_item["1,001-N"]
    assert nocturna.shift == "NOCTURNO"
    assert nocturna.codigo_sugerido == "3007 N"
    assert nocturna.capitulo_codigo == "1"


def test_preserva_el_item_original_y_recupera_el_cero(tmp_path):
    lec = leer_formulario_idu(escribir_formulario(tmp_path / "f1.xlsx"))
    originales = [i.item_pago_original for i in lec.items]
    assert "2.010" in originales          # el float 2.01 con formato 0.000
    assert "1,001-N" in originales        # el texto, tal cual venía
    assert "2,011-N" in originales


def test_las_dos_columnas_de_precio(tmp_path):
    lec = leer_formulario_idu(escribir_formulario(tmp_path / "f1.xlsx"))
    replanteo = next(i for i in lec.items if i.item_pago_original == "1.001")
    assert replanteo.cantidad == 100
    assert replanteo.precio_contractual == 1351             # col K, CON AIU
    assert replanteo.precio_contractual_sin_aiu == 1056      # col J, SIN AIU
    assert replanteo.fila_origen > 0


def test_conciliacion_contra_los_subtotales_del_excel(tmp_path):
    lec = leer_formulario_idu(escribir_formulario(tmp_path / "f1.xlsx"))
    c = lec.conciliacion
    assert c["subtotales_ok"] is True
    assert c["diferencia"] == 0
    assert c["contractual_con_aiu"] == c["subtotales_excel"]
    assert c["contractual_sin_aiu"] < c["contractual_con_aiu"]
    assert [a for a in lec.advertencias if a.tipo == "total_fila_no_concilia"] == []


def test_hoja_gemela_se_avisa_y_no_duplica_actividades(tmp_path):
    lec = leer_formulario_idu(
        escribir_formulario(tmp_path / "f1.xlsx", con_hoja_gemela=True))
    assert lec.hoja == "PROPUESTA ECONÓMICA"
    assert len(lec.items) == 5
    assert any(a.tipo == "hoja_ambigua" for a in lec.advertencias)


def test_sin_hoja_compatible_es_error_bloqueante(tmp_path):
    lec = leer_formulario_idu(escribir_formulario(tmp_path / "f1.xlsx", hoja="DATOS"))
    assert lec.items == []
    assert any("sin_hoja" in e for e in lec.errores)


def test_falta_columna_obligatoria_es_error_bloqueante(tmp_path):
    lec = leer_formulario_idu(
        escribir_formulario(tmp_path / "f1.xlsx", sin_columna_cantidad=True))
    assert any("falta_columna" in e and "cantidad" in e for e in lec.errores)
    assert lec.items == []


def test_archivo_que_no_es_excel_es_error_bloqueante(tmp_path):
    malo = tmp_path / "no-es.xlsx"
    malo.write_bytes(b"esto no es un zip")
    lec = leer_formulario_idu(malo)
    assert any("archivo_invalido" in e for e in lec.errores)
    assert lec.items == []


def test_avisa_si_la_oferta_viene_diligenciada(tmp_path):
    # En el archivo de referencia las columnas de oferta están en cero. Si algún día
    # llegan con valor, el parser NO las lee — pero avisa, en vez de callarse.
    p = escribir_formulario(tmp_path / "f1.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = wb["PROPUESTA ECONÓMICA"]
    for fila in (10, 13):                      # las dos actividades del capítulo 1
        ws.cell(row=fila, column=15).value = 1200
    wb.save(p)
    lec = leer_formulario_idu(p)
    avisos = [a for a in lec.advertencias if a.tipo == "oferta_diligenciada"]
    assert len(avisos) == 1
    assert "2" in avisos[0].detalle           # cuántas filas traen oferta
    # Y NO cambia el contractual: se sigue leyendo la columna oficial con AIU.
    assert next(i for i in lec.items
                if i.item_pago_original == "1.001").precio_contractual == 1351


def test_read_presupuesto_sigue_funcionando(tmp_path):
    # El envoltorio viejo (CLI y GUI) no cambia de firma ni de comportamiento.
    from apu_tool.dominio.presupuesto import read_presupuesto
    items = read_presupuesto(escribir_formulario(tmp_path / "f1.xlsx"),
                             hoja="PROPUESTA ECONÓMICA")
    assert len(items) == 5
    assert items[0].categoria == "1 · PRELIMINARES"
