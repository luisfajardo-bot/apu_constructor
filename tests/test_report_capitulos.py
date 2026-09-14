"""El cuadro gana una hoja por capítulo — y SOLO cuando hay capítulos."""
import openpyxl

from apu_tool.dominio.report import write_report
from apu_tool.dominio.report_categorizado import (
    resumen_por_capitulo, write_report_categorizado,
)
from apu_tool.nucleo.models import (
    AssembledApu, CostedComponent, LicitacionItem, MatchStatus,
)

HOJA = "RESUMEN POR CAPÍTULO"


def _item(cap_cod="", cap_nom="", cant=10.0, con_aiu=1351.0, sin_aiu=1056.0,
          item="1.001"):
    return LicitacionItem(
        item=item, descripcion=f"ACT {item}", unidad="M2", cantidad=cant,
        precio_contractual=con_aiu, precio_contractual_sin_aiu=sin_aiu,
        shift="DIURNO", categoria=(f"{cap_cod} · {cap_nom}" if cap_cod else ""),
        capitulo_codigo=cap_cod, capitulo_nombre=cap_nom, item_pago_original=item)


def _ens(item, costo=900.0, apu_codigo="A1"):
    comps = [CostedComponent(insumo_codigo="I1", insumo_nombre="CEMENTO", unidad="KG",
                             rendimiento=1.0, precio_unitario=costo,
                             fuente_precio="COSTO INTERNO", costo=costo)] if costo else []
    return AssembledApu(item=item, apu_codigo=apu_codigo, apu_nombre="APU",
                        unidad="M2", shift="DIURNO", componentes=comps,
                        costo_unitario=costo, status=MatchStatus.AUTO, confianza=1.0)


def _hojas(path) -> list[str]:
    wb = openpyxl.load_workbook(path)
    nombres = wb.sheetnames
    wb.close()
    return nombres


def _tabla(path, hoja):
    wb = openpyxl.load_workbook(path)
    ws = wb[hoja]
    filas = [[c.value for c in f] for f in ws.iter_rows()]
    wb.close()
    return filas[0], filas[1:]


def test_una_corrida_plana_produce_el_cuadro_de_siempre(tmp_path):
    out = tmp_path / "plano.xlsx"
    write_report([_ens(_item())], out)
    assert _hojas(out) == ["RESUMEN", "DESGLOSE", "ALERTAS", "INFO"]


def test_una_corrida_con_capitulos_gana_la_hoja(tmp_path):
    out = tmp_path / "cap.xlsx"
    write_report([_ens(_item("1", "PRELIMINARES")),
                  _ens(_item("2", "PAVIMENTOS", item="2.001"))], out)
    assert HOJA in _hojas(out)


def test_la_hoja_trae_las_columnas_pedidas(tmp_path):
    out = tmp_path / "cap.xlsx"
    write_report([_ens(_item("1", "PRELIMINARES")),
                  _ens(_item("1", "PRELIMINARES", item="1.002"), costo=0.0,
                       apu_codigo=None)], out)
    enc, filas = _tabla(out, HOJA)
    for esperado in ("Capítulo", "Nombre", "Actividades", "Con APU", "Sin APU",
                     "Contractual", "Contractual sin AIU", "Costo interno",
                     "Diferencia", "Margen %", "Cobertura", "Estado"):
        assert esperado in enc, esperado
    fila = filas[0]
    assert fila[enc.index("Capítulo")] == "1"
    assert fila[enc.index("Actividades")] == 2
    assert fila[enc.index("Sin APU")] == 1
    assert fila[enc.index("Estado")] == "Incompleto"


def test_el_gran_total_de_la_hoja_cuadra_con_los_items(tmp_path):
    out = tmp_path / "cap.xlsx"
    apus = [_ens(_item("1", "PRELIMINARES")),
            _ens(_item("2", "PAVIMENTOS", cant=4.0, item="2.001"))]
    write_report(apus, out)
    enc, filas = _tabla(out, HOJA)
    total = filas[-1]
    assert total[enc.index("Contractual")] == sum(a.contractual_total for a in apus)
    assert total[enc.index("Costo interno")] == sum(a.costo_total for a in apus)
    assert total[enc.index("Contractual sin AIU")] == sum(
        a.contractual_total_sin_aiu for a in apus)


def test_el_cuadro_categorizado_usa_la_misma_funcion(tmp_path):
    # Si los dos escritores calcularan por su cuenta, este test seria el que avisa.
    apus = [_ens(_item("1", "PRELIMINARES")),
            _ens(_item("2", "PAVIMENTOS", cant=4.0, item="2.001"))]
    esperado = resumen_por_capitulo(apus)
    out = tmp_path / "categorizado.xlsx"
    write_report_categorizado(apus, out)
    enc, filas = _tabla(out, HOJA)
    col = enc.index("Contractual")
    assert [f[col] for f in filas[:len(esperado)]] == [c["contractual"] for c in esperado]


def test_las_dos_hojas_por_capitulo_son_identicas(tmp_path):
    # `write_report` y `write_report_categorizado` escriben LA MISMA hoja: dos cuadros
    # de la misma corrida no pueden diferir en el resumen por capítulo.
    apus = [_ens(_item("1", "PRELIMINARES")),
            _ens(_item("2", "PAVIMENTOS", cant=4.0, item="2.001"), costo=0.0,
                 apu_codigo=None)]
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    write_report(apus, a)
    write_report_categorizado(apus, b)
    assert _tabla(a, HOJA) == _tabla(b, HOJA)


def test_la_hoja_resumen_trae_las_dos_bases(tmp_path):
    out = tmp_path / "cap.xlsx"
    write_report([_ens(_item("1", "PRELIMINARES"))], out)
    enc, filas = _tabla(out, "RESUMEN")
    assert "P. Contractual sin AIU" in enc
    assert "Total Contractual sin AIU" in enc
    assert filas[0][enc.index("P. Contractual sin AIU")] == 1056
    assert filas[0][enc.index("Total Contractual sin AIU")] == 10 * 1056


def test_el_detalle_categorizado_trae_las_dos_bases(tmp_path):
    out = tmp_path / "categorizado.xlsx"
    write_report_categorizado([_ens(_item("1", "PRELIMINARES"))], out)
    enc, _filas = _tabla(out, "DETALLE")
    assert "P. Contractual sin AIU" in enc
    assert "Total Contractual sin AIU" in enc
