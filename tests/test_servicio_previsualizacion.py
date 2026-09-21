"""Previsualización: lee y valida SIN escribir nada en la base."""
import json

from apu_tool.datos.almacen import Almacen
from apu_tool.nucleo.models import EntidadOrigen, LicitacionItem
from apu_tool.servicio import corridas as svc
from tests.fixtures_idu import escribir_formulario


def _almacen(tmp_path) -> Almacen:
    alm = Almacen(precios_path=tmp_path / "p.db", apus_path=tmp_path / "a.db",
                  corridas_path=tmp_path / "c.db")
    alm.init_schema()
    return alm


def test_previsualizar_no_escribe_nada(tmp_path):
    alm = _almacen(tmp_path)
    p = escribir_formulario(tmp_path / "f1.xlsx")
    svc.previsualizar(EntidadOrigen.IDU, p.read_bytes(), "f1.xlsx")
    # No hay borrador, no hay corrida a medio crear, no hay nada que limpiar.
    assert alm.corridas.listar_corridas() == []


def test_previsualizar_devuelve_capitulos_totales_y_conciliacion(tmp_path):
    p = escribir_formulario(tmp_path / "f1.xlsx")
    prev = svc.previsualizar(EntidadOrigen.IDU, p.read_bytes(), "f1.xlsx")
    assert prev["entidad"] == "IDU"
    assert prev["formato"] == "idu_formulario_1"
    assert prev["archivo"] == "f1.xlsx"
    assert prev["hoja"] == "PROPUESTA ECONÓMICA"
    assert prev["fila_encabezado"] == 5
    assert prev["parser_version"] == "idu-f1/1"
    assert prev["actividades"] == 5
    assert [c["codigo"] for c in prev["capitulos"]] == ["1", "2"]
    assert prev["capitulos"][0]["actividades"] == 2
    assert prev["capitulos"][0]["contractual"] > 0
    assert prev["capitulos"][0]["contractual_sin_aiu"] > 0
    assert prev["conciliacion"]["subtotales_ok"] is True
    assert prev["errores"] == []
    assert prev["puede_aprobar"] is True
    assert prev["requiere_confirmacion"] is True


def test_el_total_de_la_previa_es_la_suma_de_sus_capitulos(tmp_path):
    # Si el total y las filas no cuadran, el usuario aprueba una cosa distinta de la
    # que ve. Es la misma regla de "no dupliques cálculos" aplicada a la previa.
    p = escribir_formulario(tmp_path / "f1.xlsx")
    prev = svc.previsualizar(EntidadOrigen.IDU, p.read_bytes(), "f1.xlsx")
    assert prev["totales"]["contractual"] == sum(
        c["contractual"] for c in prev["capitulos"])
    assert prev["totales"]["contractual_sin_aiu"] == sum(
        c["contractual_sin_aiu"] for c in prev["capitulos"])
    # Y con lo que el parser concilió contra el Excel.
    assert prev["totales"]["contractual"] == prev["conciliacion"]["contractual_con_aiu"]


def test_previsualizar_con_error_no_deja_aprobar(tmp_path):
    p = escribir_formulario(tmp_path / "f1.xlsx", sin_columna_cantidad=True)
    prev = svc.previsualizar(EntidadOrigen.IDU, p.read_bytes(), "f1.xlsx")
    assert prev["errores"] != []
    assert prev["puede_aprobar"] is False
    assert prev["capitulos"] == []


def test_previsualizar_con_advertencias_deja_aprobar_pero_las_muestra(tmp_path):
    p = escribir_formulario(tmp_path / "f1.xlsx", con_hoja_gemela=True)
    prev = svc.previsualizar(EntidadOrigen.IDU, p.read_bytes(), "f1.xlsx")
    assert prev["puede_aprobar"] is True
    assert any(a["tipo"] == "hoja_ambigua" for a in prev["advertencias"])


def test_previsualizar_limita_las_filas_senaladas(tmp_path):
    p = escribir_formulario(tmp_path / "f1.xlsx")
    prev = svc.previsualizar(EntidadOrigen.IDU, p.read_bytes(), "f1.xlsx")
    assert len(prev["filas_senaladas"]) <= svc.MAX_FILAS_SENALADAS
    # Solo las que apuntan a una fila concreta: un aviso global no es "una fila".
    assert all(f["fila"] > 0 for f in prev["filas_senaladas"])


def test_previsualizar_no_devuelve_las_actividades_una_por_una(tmp_path):
    # La previa valida la ESTRUCTURA; la tabla de la corrida muestra las actividades.
    # Mandar 1939 filas para que el usuario mire 3 es medio mega por nada.
    p = escribir_formulario(tmp_path / "f1.xlsx")
    prev = svc.previsualizar(EntidadOrigen.IDU, p.read_bytes(), "f1.xlsx")
    assert "items" not in prev


def test_previsualizar_de_una_entidad_generica_no_trae_capitulos(tmp_path):
    import openpyxl
    p = tmp_path / "plana.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ITEM", "DESCRIPCION", "UNIDAD", "CANTIDAD", "PRECIO", "TURNO"])
    ws.append(["1", "EXCAVACION MANUAL", "M3", 10, 45000, "DIURNO"])
    wb.save(p)
    prev = svc.previsualizar(EntidadOrigen.INVIAS, p.read_bytes(), "plana.xlsx")
    assert prev["capitulos"] == []
    assert prev["actividades"] == 1
    assert prev["puede_aprobar"] is True
    assert prev["requiere_confirmacion"] is False
    assert prev["formato"] == "generico"


def test_crear_corrida_guarda_el_origen(tmp_path):
    alm = _almacen(tmp_path)
    p = escribir_formulario(tmp_path / "f1.xlsx")
    lectura = svc.leer_para_corrida(EntidadOrigen.IDU, p.read_bytes(), "f1.xlsx")
    cid = svc.crear_corrida_encolada(
        alm, "f1.xlsx", lectura.items, "DIURNO", False,
        origen=svc.origen_de(EntidadOrigen.IDU, lectura, "f1.xlsx", "yo@test.co"))
    origen = json.loads(alm.corridas.get_origen(cid))
    assert origen["entidad"] == "IDU"
    assert origen["formato"] == "idu_formulario_1"
    assert origen["hoja"] == "PROPUESTA ECONÓMICA"
    assert origen["parser_version"] == "idu-f1/1"
    assert origen["estructura_confirmada"] is True
    assert origen["confirmada_por"] == "yo@test.co"
    assert origen["capitulos"] == 2
    assert origen["actividades"] == 5
    assert origen["conciliacion"]["contractual_con_aiu"] > 0
    assert origen["importada_en"]


def test_crear_corrida_sin_origen_sigue_funcionando(tmp_path):
    # El camino de la CLI/GUI y el de las entidades sin lector especializado.
    alm = _almacen(tmp_path)
    items = [LicitacionItem(item="1", descripcion="X", unidad="M2", cantidad=1.0,
                            precio_contractual=100.0, shift="DIURNO")]
    cid = svc.crear_corrida_encolada(alm, "plana.xlsx", items, "DIURNO", False)
    assert alm.corridas.get_origen(cid) is None
    assert alm.corridas.get_corrida(cid).origen is None


def test_el_origen_cuenta_las_advertencias_sin_guardar_su_texto(tmp_path):
    # El detalle ya lo mostró la previa y el usuario ya lo confirmó; lo que importa
    # después es el rastro de que las hubo. Además `detalle` lleva montos en texto.
    p = escribir_formulario(tmp_path / "f1.xlsx", con_hoja_gemela=True)
    lectura = svc.leer_para_corrida(EntidadOrigen.IDU, p.read_bytes(), "f1.xlsx")
    origen = svc.origen_de(EntidadOrigen.IDU, lectura, "f1.xlsx", "yo@test.co")
    # El fixture también trae el encabezado repetido a mitad: se cuentan los dos.
    assert origen["advertencias"] == {"hoja_ambigua": 1, "encabezado_repetido": 1}


def test_una_actividad_huerfana_si_aparece_como_grupo(tmp_path):
    # Con capítulos declarados, una actividad antes del primero es una ANOMALÍA que el
    # usuario tiene que ver — al revés de la lista plana, donde no hay nada que mostrar.
    import openpyxl
    p = escribir_formulario(tmp_path / "f1.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = wb["PROPUESTA ECONÓMICA"]
    ws.insert_rows(6)                       # justo después del encabezado
    for col, val in ((3, 4444), (4, "0.900"), (7, "HUERFANA"), (8, "M2"),
                     (9, 10), (10, 100), (11, 128)):
        ws.cell(row=6, column=col).value = val
    wb.save(p)
    prev = svc.previsualizar(EntidadOrigen.IDU, p.read_bytes(), "f1.xlsx")
    assert any(c["codigo"] == "" and c["nombre"] == "(sin capítulo)"
               for c in prev["capitulos"])
    assert any(a["tipo"] == "actividad_sin_capitulo" for a in prev["advertencias"])
